import math
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class DocumentParseError(ValueError):
    """Raised when text cannot be extracted from an uploaded file.

    Routes translate this into a 400 so a corrupt or password-protected file is
    rejected at upload time rather than failing later inside the background job.
    """


def extract_text_from_pdf(file_path: Path) -> tuple[str, int]:
    """Return ``(text, page_count)`` for a PDF."""
    try:
        reader = PdfReader(str(file_path))
        if reader.is_encrypted:
            raise DocumentParseError(
                "This PDF is password protected, so its text cannot be read. "
                "Please upload an unprotected copy."
            )
        # Read every page. The previous implementation returned inside this loop
        # and so only ever captured the first page.
        parts = [(page.extract_text() or "") for page in reader.pages]
    except DocumentParseError:
        raise
    except Exception as exc:  # pypdf raises a variety of low-level errors
        raise DocumentParseError(f"Could not read the PDF: {exc}") from exc

    return "\n".join(parts).strip(), len(reader.pages)


def extract_text_from_docx(file_path: Path) -> str:
    """Return the text of a DOCX, including table cells."""
    try:
        document = Document(str(file_path))
    except Exception as exc:
        raise DocumentParseError(f"Could not read the DOCX file: {exc}") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return "\n".join(blocks).strip()


def extract_text_from_txt(file_path: Path) -> str:
    """Return the text of a plain text file, tolerating non-UTF-8 encodings."""
    raw = file_path.read_bytes()
    try:
        return raw.decode("utf-8").strip()
    except UnicodeDecodeError:
        return raw.decode("latin-1").strip()


def extract_text(file_path: Path) -> str:
    """Extract the raw text of a supported document, dispatching on suffix."""
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(file_path)[0]
    if suffix == ".docx":
        return extract_text_from_docx(file_path)
    if suffix == ".txt":
        return extract_text_from_txt(file_path)

    raise DocumentParseError(f"Unsupported file type '{suffix}'.")


def extract_contract(file_path: Path) -> tuple[str, int, int]:
    """Return ``(text, page_count, word_count)`` for a supported document.

    Raises DocumentParseError if the file cannot be read, or if it yields no
    usable text (for example a scanned PDF with no text layer), since a contract
    with no text cannot be analysed.
    """
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        text, page_count = extract_text_from_pdf(file_path)
    elif suffix == ".docx":
        text = extract_text_from_docx(file_path)
        # python-docx has no notion of rendered pages - page count depends on
        # layout, not the document model. This is a rough estimate at ~500 words
        # per page, and is only used for display.
        page_count = max(1, math.ceil(len(text.split()) / 500))
    elif suffix == ".txt":
        text = extract_text_from_txt(file_path)
        page_count = 1
    else:
        raise DocumentParseError(f"Unsupported file type '{suffix}'.")

    if not text.strip():
        raise DocumentParseError(
            "No text could be extracted from this file. If it is a scanned "
            "document it may need to be run through OCR first."
        )

    return text, page_count, len(text.split())
