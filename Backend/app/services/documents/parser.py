"""Text extraction from PDF, DOCX and TXT bytes. Synchronous and CPU-bound:
call through ``extract_document``, which runs it in a worker thread."""

import io
import math
from dataclasses import dataclass

import pypdfium2 as pdfium
from docx import Document


class DocumentParseError(ValueError):
    """The file cannot be read. Shown to the user as a 400 at upload time."""


@dataclass(slots=True)
class ParsedText:
    text: str
    page_count: int


def parse_pdf(data: bytes) -> ParsedText:
    """PDF text via PDFium (the engine inside Chrome).

    PDFium keeps words whole where pypdf often split them ("C ontract",
    "promis e"), which hurt both search and quote checking.
    """
    try:
        document = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        if "password" in str(exc).lower():
            raise DocumentParseError(
                "This PDF is password protected. Please upload an unprotected copy."
            ) from exc
        raise DocumentParseError(
            "We couldn't open this PDF. It may be damaged. Try saving or exporting it again."
        ) from exc

    try:
        pages = [page.get_textpage().get_text_range() for page in document]
    finally:
        document.close()

    text = "\n".join(pages).replace("\r\n", "\n").replace("\r", "\n")
    return ParsedText(text.strip(), len(pages))


def parse_docx(data: bytes) -> ParsedText:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types
        raise DocumentParseError(
            "We couldn't open this Word file. Try saving it again as .docx, or upload a PDF."
        ) from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            if cells := [cell.text.strip() for cell in row.cells if cell.text.strip()]:
                blocks.append(" | ".join(cells))

    text = "\n".join(blocks).strip()
    # DOCX has no fixed pages; ~500 words a page is a display estimate only.
    return ParsedText(text, max(1, math.ceil(len(text.split()) / 500)))


def parse_txt(data: bytes) -> ParsedText:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return ParsedText(text.strip(), 1)


PARSERS = {".pdf": parse_pdf, ".docx": parse_docx, ".txt": parse_txt}
